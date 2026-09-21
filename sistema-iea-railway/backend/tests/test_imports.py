"""All fixtures are fictitious. Reuses the disposable-database guard."""
import io
import json
import unittest
from unittest.mock import patch
import test_institution_api as base
from app import import_service as service
from openpyxl import Workbook, load_workbook


class ImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        base.InstitutionApiTests.setUpClass()
        cls.client = base.InstitutionApiTests.client

    def setUp(self):
        base.InstitutionApiTests.setUp(self)
        base.sql("INSERT INTO sedes(id,nombre) VALUES (2,'Avellaneda')")
        base.sql("INSERT INTO asignaciones(id,catedra_id,docente_id,cuatrimestre_id,sede_id,modalidad,dia,hora_inicio,hora_fin,recibe_alumnos_presenciales,modificada,comision) VALUES (1,1,1,1,1,'presencial','Lunes','18:00','19:30',TRUE,TRUE,'A'),(2,2,1,1,2,'presencial','Martes','18:00','19:30',FALSE,FALSE,'B'),(3,1,1,2,1,'presencial','Viernes','18:00','19:30',FALSE,FALSE,'A')")
        base.sql("UPDATE inscripciones SET sede_referencia='Caballito',modalidad_alumno='presencial',turno='Noche'")
        for table in service.TABLES:
            base.sql(f"SELECT setval(pg_get_serial_sequence('{table}','id'),coalesce(max(id),1),max(id) IS NOT NULL) FROM {table}")

    def state(self):
        with base.SessionLocal() as db: return service.digest(service.read_state(db))

    def file(self, rows=None, title='Datos'):
        if rows is None and hasattr(self,'default_file'): return self.default_file
        workbook=Workbook(); sheet=workbook.active; sheet.title=title
        for row in rows or [
            ['asignacion_id','codigo_materia','materia','sede','comision','dia','hora_inicio','hora_fin','docente','modalidad','link_meet','recibe_alumnos_presenciales'],
            [1,'c.1','Materia de prueba','Caballito','A','Miércoles','19:00','20:00','Docente Ficticio','presencial','https://example.invalid/clase','si'],
        ]: sheet.append(row)
        output=io.BytesIO(); workbook.save(output)
        content=output.getvalue()
        if rows is None: self.default_file=content
        return content

    def options(self,kind='horarios',campus=1,period=1,mode='actualizar',formato='auto'):
        data={'tipo':kind,'sede_id':campus,'modo':mode,'formato':formato}
        if period is not None: data['cuatrimestre_id']=period
        return data

    def preview(self,content=None,options=None,mapping=None):
        response=self.client.post('/api/importaciones/vista-previa',params=options or self.options(),
            files={'file':('prueba.xlsx',content or self.file())},data={'mapeo':json.dumps(mapping or {})})
        self.assertEqual(response.status_code,200,response.text[:500])
        return response.json()

    def apply(self,preview,content=None,options=None,mapping=None):
        return self.client.post('/api/importaciones/aplicar',params=options or self.options(),
            files={'file':('prueba.xlsx',content or self.file())},data={'token':preview.get('token') or 'invalid','mapeo':json.dumps(mapping or {})})

    def restore(self,history_id):
        response=self.client.post(f'/api/importaciones/historial/{history_id}/vista-previa')
        self.assertEqual(response.status_code,200,response.text[:200])
        preview=response.json()
        return self.client.post(f'/api/importaciones/historial/{history_id}/restaurar',json={'token':preview['token']}),preview

    def test_preview_is_read_only_including_aliases_and_history(self):
        before=self.state(); preview=self.preview()
        self.assertTrue(preview['puede_aplicar'])
        self.assertEqual(preview['resumen']['modificaciones'],1)
        self.assertEqual(before,self.state())
        self.assertEqual(base.sql('SELECT count(*) FROM importaciones_historial')[0][0],0)

    def test_unknown_teacher_name_does_not_create_an_unidentified_duplicate(self):
        before=self.state()
        content=self.file([['codigo','dia','hora','sede','docente'],['c.1','Jueves','18:00','Caballito','Nombre no identificado']])
        preview=self.preview(content)
        self.assertFalse(preview['puede_aplicar'])
        self.assertIn('docente_id',preview['errores'][0]['mensaje'])
        self.assertEqual(before,self.state())

    def test_teacher_public_id_resolves_and_conflicting_document_blocks_import(self):
        content=self.file([['codigo','dia','hora','sede','docente_id','docente_dni'],['c.1','Jueves','18:00','Caballito','DOC-000001','99.000.001']])
        preview=self.preview(content)
        self.assertTrue(preview['puede_aplicar'],preview['errores'])
        self.assertEqual(self.apply(preview,content).status_code,200)
        self.assertEqual(base.sql('SELECT count(*) FROM docentes')[0][0],1)
        conflict=self.file([['codigo','dia','hora','sede','docente_id','docente_dni'],['c.1','Jueves','18:00','Caballito','DOC-000001','99700002']])
        self.assertFalse(self.preview(conflict)['puede_aplicar'])

    def test_partial_replace_preserves_other_campus_and_period(self):
        options=self.options(mode='reemplazar')
        before=base.sql('SELECT * FROM asignaciones WHERE id IN (2,3) ORDER BY id')
        response=self.apply(self.preview(options=options),options=options)
        self.assertEqual(response.status_code,200,response.text[:300])
        self.assertEqual(before,base.sql('SELECT * FROM asignaciones WHERE id IN (2,3) ORDER BY id'))
        self.assertEqual(base.sql('SELECT hora_inicio FROM asignaciones WHERE id=1')[0][0],'19:00')

    def test_update_preserves_absent_rows_while_replace_proposes_deletions(self):
        base.sql("INSERT INTO asignaciones(catedra_id,cuatrimestre_id,sede_id,modalidad,dia,hora_inicio,comision) VALUES (2,1,1,'presencial','Jueves','18:00','C')")
        self.assertEqual(self.preview()['resumen']['bajas'],0)
        self.assertEqual(self.preview(options=self.options(mode='reemplazar'))['resumen']['bajas'],1)

    def test_reimport_is_idempotent_and_preserves_ids(self):
        self.assertEqual(self.apply(self.preview()).status_code,200)
        before=self.state(); second=self.preview()
        self.assertEqual(second['resumen']['altas'],0)
        self.assertEqual(second['resumen']['modificaciones'],0)
        response=self.apply(second)
        self.assertTrue(response.json()['sin_cambios'])
        self.assertEqual(before,self.state())

    def test_invalid_empty_and_unknown_subject_files_never_replace(self):
        before=self.state(); options=self.options(mode='reemplazar')
        for rows in ([['x']], [['codigo','dia','hora','sede'],['c.404','Lunes','18:00','Caballito']],
                     [['codigo','dia','hora','sede'],['c.1','Día Inventado','25:99','Caballito']]):
            with self.subTest(rows=rows):
                content=self.file(rows); preview=self.preview(content,options)
                self.assertFalse(preview['puede_aplicar']); self.assertTrue(preview['errores'])
                self.assertGreaterEqual(self.apply(preview,content,options).status_code,400)
                self.assertEqual(before,self.state())

    def test_changed_file_or_scope_invalidates_preview(self):
        preview=self.preview(); before=self.state()
        changed=self.file([['codigo','dia','hora','sede'],['c.2','Lunes','17:00','Caballito']])
        self.assertEqual(self.apply(preview,changed).status_code,409)
        self.assertEqual(self.apply(preview,options=self.options(mode='reemplazar')).status_code,409)
        self.assertEqual(before,self.state())

    def test_stale_preview_preserves_intervening_change(self):
        preview=self.preview()
        base.sql("UPDATE asignaciones SET hora_inicio='17:00' WHERE id=1")
        before=self.state()
        self.assertEqual(self.apply(preview).status_code,409)
        self.assertEqual(before,self.state())

    def test_no_business_writes_when_backup_is_unavailable(self):
        preview=self.preview(); before=self.state()
        base.sql('ALTER TABLE importaciones_historial RENAME TO historial_temporal_prueba')
        try:
            self.assertEqual(self.apply(preview).status_code,500)
            self.assertEqual(before,self.state())
        finally: base.sql('ALTER TABLE historial_temporal_prueba RENAME TO importaciones_historial')

    def test_failure_mid_import_rolls_back_everything(self):
        content=self.file([['codigo','dia','hora','sede','docente','docente_dni'],['c.1','Jueves','18:00','Caballito','Nueva Persona Ficticia','99700002']])
        preview=self.preview(content); before=self.state()
        real=service.write_row; calls=[]
        def fail_second(*args):
            calls.append(1)
            if len(calls)==2: raise RuntimeError('Falla simulada después de la primera escritura')
            return real(*args)
        with patch.object(service,'write_row',side_effect=fail_second):
            self.assertEqual(self.apply(preview,content).status_code,500)
        self.assertEqual(before,self.state())
        self.assertEqual(base.sql('SELECT count(*) FROM importaciones_historial')[0][0],0)

    def test_restore_recovers_all_affected_fields_and_new_entities(self):
        before=self.state()
        content=self.file([['asignacion_id','codigo','dia','hora','hora fin','sede','docente','modalidad','link meet','recibe alumnos presenciales','docente_dni'],
                          [1,'c.1','Jueves','20:00','21:00','Caballito','Nueva Persona Ficticia','presencial','https://example.invalid/nuevo','no','99700002']])
        result=self.apply(self.preview(content),content)
        self.assertEqual(result.status_code,200,result.text[:400])
        restored,preview=self.restore(result.json()['historial_id'])
        self.assertEqual(restored.status_code,200,restored.text[:400])
        self.assertEqual(before,self.state())

    def test_restore_rejects_overwriting_later_edits(self):
        result=self.apply(self.preview()).json()
        base.sql("UPDATE catedras SET link_meet='https://example.invalid/posterior' WHERE id=1")
        before=self.state()
        restored,preview=self.restore(result['historial_id'])
        self.assertFalse(preview['puede_restaurar'])
        self.assertEqual(restored.status_code,409)
        self.assertEqual(before,self.state())

    def test_enrollment_decimal_dni_update_replace_and_restore(self):
        content=self.file([['codigo_materia','dni','nombre','apellido','sede','modalidad'],['c.1','98000001.0','Alumno','Ficticio','Caballito','presencial']])
        options=self.options(kind='inscripciones',mode='reemplazar')
        before=self.state(); preview=self.preview(content,options)
        self.assertTrue(preview['puede_aplicar'],preview['errores'])
        self.assertEqual(preview['resumen']['altas'],0)
        self.assertEqual(preview['resumen']['bajas'],110)
        result=self.apply(preview,content,options)
        self.assertEqual(result.status_code,200,result.text[:300])
        self.assertEqual(base.sql('SELECT count(*) FROM inscripciones')[0][0],1)
        self.assertEqual(base.sql('SELECT count(*) FROM alumnos')[0][0],101)
        self.assertEqual(self.restore(result.json()['historial_id'])[0].status_code,200)
        self.assertEqual(before,self.state())

    def test_scoped_enrollment_import_preserves_unclassified_and_other_campus(self):
        base.sql("UPDATE inscripciones SET sede_referencia='Avellaneda' WHERE catedra_id=2")
        base.sql("UPDATE inscripciones SET sede_referencia=NULL WHERE alumno_id=2")
        protected=base.sql("SELECT * FROM inscripciones WHERE sede_referencia IS NULL OR sede_referencia='Avellaneda' ORDER BY id")
        content=self.file([['codigo_materia','dni','nombre','sede'],['c.1','98000001','Alumno','Caballito']])
        options=self.options(kind='inscripciones',mode='reemplazar')
        result=self.apply(self.preview(content,options),content,options)
        self.assertEqual(result.status_code,200,result.text[:200])
        self.assertEqual(protected,base.sql("SELECT * FROM inscripciones WHERE sede_referencia IS NULL OR sede_referencia='Avellaneda' ORDER BY id"))

    def test_plan_partial_replace_and_restore(self):
        base.sql("INSERT INTO plan_carrera(sede,carrera,anno,codigo_catedra,nombre_catedra) VALUES ('Caballito','Carrera prueba','1ER AÑO','c.1','Primera'),('Avellaneda','Carrera prueba','1ER AÑO','c.2','Segunda')")
        content=self.file([['codigo_materia','materia','carrera','sede','anno'],['c.1','Primera actualizada','Carrera prueba','Caballito','1ER AÑO']])
        options=self.options(kind='plan',period=None,mode='reemplazar')
        before=self.state()
        result=self.apply(self.preview(content,options),content,options)
        self.assertEqual(result.status_code,200,result.text[:300])
        self.assertEqual(base.sql("SELECT nombre_catedra FROM plan_carrera WHERE sede='Avellaneda'")[0][0],'Segunda')
        self.assertEqual(self.restore(result.json()['historial_id'])[0].status_code,200)
        self.assertEqual(before,self.state())

    def test_iea_and_another_institution_headers_use_same_pipeline(self):
        iea=self.file([['CODIGO','MATERIA','DIA','HORA INICIO','HORA FIN','SEDE','DOCENTE'],[1,'Prueba','Lunes','18:00','19:00','Caballito','Docente Ficticio']])
        other=self.file([['Subject','Campus name','Weekday','Begin','Teacher'],['c.1','Caballito','Lunes','18:00','Docente Ficticio']])
        mapping={'codigo_materia':'Subject','sede':'Campus name','dia':'Weekday','hora_inicio':'Begin','docente':'Teacher'}
        self.assertTrue(self.preview(iea)['puede_aplicar'])
        custom=self.preview(other,mapping=mapping)
        self.assertTrue(custom['puede_aplicar'],custom['errores'])
        self.assertEqual(self.apply(custom,other,mapping=mapping).status_code,200)

    def test_legacy_direct_imports_cannot_bypass_preview(self):
        before=self.state()
        for path in ('horarios-aplicar','alumnos','plan-carrera'):
            response=self.client.post('/api/importar/'+path,files={'file':('prueba.xlsx',self.file())})
            self.assertEqual(response.status_code,409)
        self.assertEqual(before,self.state())

    def test_exported_ids_edit_one_session_without_affecting_other_campus(self):
        response=self.client.get('/api/exportar/planilla-trabajo?cuatrimestre_id=1&sede=Caballito')
        self.assertEqual(response.status_code,200)
        workbook=load_workbook(io.BytesIO(response.content)); sheet=workbook.active
        headers={cell.value:cell.column for cell in sheet[1]}
        # Export includes an empty planning row for local demand; finish or remove it before importing.
        for row in range(sheet.max_row,1,-1):
            if not sheet.cell(row,headers['ASIGNACION_ID']).value: sheet.delete_rows(row)
        self.assertEqual(sheet.max_row,2)
        self.assertEqual(sheet.cell(2,headers['ASIGNACION_ID']).value,1)
        self.assertEqual(sheet.cell(2,headers['DOCENTE_ID']).value,'DOC-000001')
        self.assertEqual(sheet.cell(2,headers['CODIGO']).value,'c.1')
        self.assertEqual(sheet.cell(2,headers['SEDE']).value,'Caballito')
        sheet.cell(2,headers['HORA INICIO'],'18:30')
        output=io.BytesIO(); workbook.save(output); content=output.getvalue()
        protected=base.sql('SELECT * FROM asignaciones WHERE id IN (2,3) ORDER BY id')
        preview=self.preview(content)
        self.assertTrue(preview['puede_aplicar'],preview['errores'])
        self.assertEqual(preview['resumen']['altas'],0)
        self.assertEqual(self.apply(preview,content).status_code,200)
        self.assertEqual(base.sql('SELECT hora_inicio FROM asignaciones WHERE id=1')[0][0],'18:30')
        self.assertEqual(protected,base.sql('SELECT * FROM asignaciones WHERE id IN (2,3) ORDER BY id'))

    def test_one_commission_can_have_two_sessions_without_duplicates(self):
        rows=[['codigo','sede','dia','hora','comision','modalidad'],
              ['c.1','Caballito','Lunes','10:00','C','presencial'],
              ['c.1','Caballito','Jueves','10:00','C','presencial']]
        content=self.file(rows+[rows[1]])
        preview=self.preview(content)
        self.assertTrue(preview['puede_aplicar'],preview['errores'])
        self.assertEqual(preview['resumen']['altas'],2)
        self.assertEqual(self.apply(preview,content).status_code,200)
        self.assertEqual(self.preview(content)['resumen']['altas'],0)
        self.assertEqual(base.sql("SELECT count(*) FROM asignaciones WHERE comision='C'")[0][0],2)

    def test_classic_iea_plan_and_enrollment_adapters(self):
        plan=self.file([['','TECNICATURA DE PRUEBA'],['','','1ER AÑO',1,'Materia de prueba','','Lunes','10:00']],title='Caballito')
        options=self.options(kind='plan',period=None,formato='iea')
        preview=self.preview(plan,options)
        self.assertTrue(preview['puede_aplicar'],preview['errores'])
        self.assertEqual(self.apply(preview,plan,options).status_code,200)
        self.assertEqual(base.sql('SELECT codigo_catedra,anno FROM plan_carrera')[0],('c.1','1ER AÑO'))
        content=self.file([['N','Apellido y nombre','Documento','Asignatura','Curso'],
                           [1,'Persona Ficticia',97000001,'c.1 - NOCHE','TECNICATURA (CABALLITO)'],
                           [2,'Otra Persona',97000002,'EDI','TECNICATURA (CABALLITO)']])
        options=self.options(kind='inscripciones',formato='iea')
        preview=self.preview(content,options)
        self.assertTrue(preview['puede_aplicar'],preview['errores'])
        self.assertEqual(self.apply(preview,content,options).status_code,200)
        self.assertEqual(base.sql('SELECT count(*) FROM inscripciones WHERE es_edi=TRUE')[0][0],1)

    def test_custom_institution_csv_with_its_own_campus_and_codes(self):
        profile=base.load_institution(base.PROFILES/'institucion-ejemplo.json')
        base.sql("UPDATE sedes SET nombre='Campus Norte' WHERE id=1")
        base.sql("INSERT INTO catedras(codigo,nombre) VALUES ('ALG101','Álgebra de prueba')")
        content='Subject;Campus name;Weekday;Begin;Mode\nALG101;Campus Norte;Martes;11:15;presencial\n'.encode()
        mapping={'codigo_materia':'Subject','sede':'Campus name','dia':'Weekday','hora_inicio':'Begin','modalidad':'Mode'}
        options=self.options(formato='estandar')
        with patch.object(service,'INSTITUCION',profile):
            before=self.state()
            response=self.client.post('/api/importaciones/vista-previa',params=options,files={'file':('otra.csv',content)},data={'mapeo':json.dumps(mapping)})
            self.assertEqual(response.status_code,200)
            preview=response.json(); self.assertTrue(preview['puede_aplicar'],preview['errores'])
            result=self.client.post('/api/importaciones/aplicar',params=options,files={'file':('otra.csv',content)},data={'mapeo':json.dumps(mapping),'token':preview['token']})
            self.assertEqual(result.status_code,200,result.text[:300])
            self.assertEqual(self.restore(result.json()['historial_id'])[0].status_code,200)
            self.assertEqual(before,self.state())

    def test_restore_deleted_assignment_with_same_id_and_flags(self):
        base.sql("INSERT INTO asignaciones(id,catedra_id,cuatrimestre_id,sede_id,modalidad,dia,hora_inicio,recibe_alumnos_presenciales,modificada,comision,carrera,turno) VALUES (4,2,1,1,'presencial','Jueves','11:00',TRUE,TRUE,'X','Carrera anterior','Mañana')")
        before=self.state(); options=self.options(mode='reemplazar')
        result=self.apply(self.preview(options=options),options=options)
        self.assertEqual(result.status_code,200,result.text[:300])
        self.assertEqual(base.sql('SELECT count(*) FROM asignaciones WHERE id=4')[0][0],0)
        self.assertEqual(self.restore(result.json()['historial_id'])[0].status_code,200)
        self.assertEqual(before,self.state())

    def test_failed_restoration_rolls_back_and_can_be_retried(self):
        result=self.apply(self.preview()).json(); before=self.state()
        history_count=base.sql('SELECT count(*) FROM importaciones_historial')[0][0]
        with patch.object(service,'write_row',side_effect=RuntimeError('Falla de recuperación simulada')):
            self.assertEqual(self.restore(result['historial_id'])[0].status_code,500)
        self.assertEqual(before,self.state())
        self.assertEqual(history_count,base.sql('SELECT count(*) FROM importaciones_historial')[0][0])
        self.assertEqual(self.restore(result['historial_id'])[0].status_code,200)

    def test_restore_preserves_new_references_and_corrupt_history_blocks(self):
        content=self.file([['codigo','sede','dia','hora','docente','docente_dni'],['c.1','Caballito','Jueves','15:00','Nuevo Docente Ficticio','99700002']])
        result=self.apply(self.preview(content),content).json()
        teacher_id=base.sql("SELECT id FROM docentes WHERE nombre='Nuevo Docente Ficticio'")[0][0]
        base.sql("INSERT INTO docente_alias(alias,docente_id) VALUES ('Alias posterior',:id)",{'id':teacher_id})
        before=self.state()
        self.assertEqual(self.restore(result['historial_id'])[0].status_code,409)
        self.assertEqual(before,self.state())
        base.sql("UPDATE importaciones_historial SET datos='invalid' WHERE id=:id",{'id':result['historial_id']})
        self.assertEqual(self.client.post(f"/api/importaciones/historial/{result['historial_id']}/vista-previa").status_code,422)

    def test_missing_scope_formulas_and_incomplete_manual_mapping_block(self):
        response=self.client.post('/api/importaciones/vista-previa?tipo=horarios',files={'file':('prueba.xlsx',self.file())})
        self.assertEqual(response.status_code,422)
        formulas=self.file([['codigo','sede','dia','hora','docente'],['c.1','Caballito','Lunes','12:00','=A1']])
        self.assertFalse(self.preview(formulas)['puede_aplicar'])
        # Manual mapping must not silently fall back to positional IEA interpretation.
        positional=self.file([[1,'Materia','Lunes','12:00','Caballito']])
        self.assertFalse(self.preview(positional,mapping={'codigo_materia':'Missing'})['puede_aplicar'])

    def test_migration_upgrades_old_schema_without_losing_assignments(self):
        expected=base.sql('SELECT id,catedra_id,dia,recibe_alumnos_presenciales,modificada FROM asignaciones ORDER BY id')
        base.sql('ALTER TABLE asignaciones DROP COLUMN comision, DROP COLUMN carrera, DROP COLUMN turno')
        base.sql('DROP TABLE importaciones_historial')
        with base.SessionLocal() as db: base.m.run_migration(db)
        self.assertEqual(expected,base.sql('SELECT id,catedra_id,dia,recibe_alumnos_presenciales,modificada FROM asignaciones ORDER BY id'))
        self.assertTrue(self.preview()['puede_aplicar'])

    def test_legacy_backup_restoration_keeps_complete_fields_and_rolls_back_failure(self):
        original=self.state()
        with base.SessionLocal() as db: history_id=base.m.crear_respaldo(db,1)
        base.sql("UPDATE asignaciones SET modificada=FALSE,comision='Cambiada' WHERE id=1")
        before=self.state()
        with patch.object(base.m,'crear_respaldo',side_effect=base.m.HTTPException(500,'Falla simulada de respaldo')):
            response=self.client.post(f'/api/respaldos/{history_id}/restaurar')
            self.assertEqual(response.status_code,500)
        self.assertEqual(before,self.state())
        self.assertEqual(self.client.post(f'/api/respaldos/{history_id}/restaurar').status_code,200)
        self.assertEqual(original,self.state())


if __name__=='__main__': unittest.main(verbosity=2)
