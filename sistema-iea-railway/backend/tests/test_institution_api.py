"""Integration tests: destructive fixtures, exclusively for a disposable local DB.

Set TEST_DATABASE_DISPOSABLE=yes and DATABASE_URL to the local test database.
The runner used during development is documented in docs/ESTANDARIZACION.md.
"""
import os
import unittest
from urllib.parse import urlparse

if os.environ.get('TEST_DATABASE_DISPOSABLE') != 'yes':
    raise unittest.SkipTest('Requires an explicitly disposable local database')
if urlparse(os.environ.get('DATABASE_URL', '')).hostname not in ('localhost', '127.0.0.1'):
    raise RuntimeError('These destructive tests require a local disposable database')

import asyncio
import io
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
from sqlalchemy import text
from openpyxl import Workbook, load_workbook
from app import main as m
from app.database import SessionLocal
from app.institution import load_institution

PROFILES = Path(__file__).resolve().parents[1] / 'app' / 'profiles'


def sql(query, params=None):
    with SessionLocal() as db:
        result = db.execute(text(query), params or {})
        rows = result.fetchall() if result.returns_rows else []
        db.commit()
        return rows


class InstitutionApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with SessionLocal() as db:
            m.run_migration(db)
        cls.client = TestClient(m.app, raise_server_exceptions=True)

    def setUp(self):
        tables = sql("SELECT tablename FROM pg_tables WHERE schemaname='public'")
        sql('TRUNCATE ' + ','.join('"' + row[0] + '"' for row in tables) + ' RESTART IDENTITY CASCADE')
        sql("INSERT INTO sedes(id,nombre) VALUES (1,'Caballito')")
        sql("INSERT INTO cuatrimestres(id,nombre,anio,numero,activo) VALUES (1,'Prueba 2026',2026,1,true),(2,'Prueba 2027',2027,1,false)")
        sql("INSERT INTO catedras(id,codigo,nombre) VALUES (1,'c.1','Materia de prueba'),(2,'c.2','Materia de prueba 2')")
        sql("INSERT INTO docentes(id,dni,nombre,apellido) VALUES (1,'99000001','Docente','Ficticio')")
        sql("INSERT INTO alumnos(id,dni,nombre,apellido) SELECT i,(98000000+i)::text,'Alumno ficticio',i::text FROM generate_series(1,101) i")
        sql("INSERT INTO inscripciones(alumno_id,catedra_id,cuatrimestre_id) SELECT i,1,1 FROM generate_series(1,101) i")
        sql("INSERT INTO inscripciones(alumno_id,catedra_id,cuatrimestre_id) SELECT i,2,1 FROM generate_series(1,10) i")
        sql("INSERT INTO catedra_dictado(catedra_id,cuatrimestre_id,se_dicta) VALUES (1,1,true),(2,1,true)")

    def get(self, route):
        response = self.client.get(route)
        self.assertEqual(response.status_code, 200, response.text[:200])
        return response.json()

    def profile(self, name):
        return patch.object(m, 'INSTITUCION', load_institution(PROFILES / name))

    def test_iea_profile_retains_opening_and_staffing(self):
        with self.profile('iea.json'):
            criteria = self.get('/api/catedras/criterio-apertura?cuatrimestre_id=1')
            self.assertEqual({r['codigo']: r['docentes_sugeridos'] for r in criteria['abrir']}, {'c.1': 2, 'c.2': 1})

    def test_custom_rules_agree_in_catalog_criteria_checklist_and_excel(self):
        with self.profile('institucion-ejemplo.json'):
            self.assertEqual(self.get('/api/institucion')['minimo_inscriptos_apertura'], 15)
            catalog = self.get('/api/catedras?cuatrimestre_id=1')
            criteria = self.get('/api/catedras/criterio-apertura?cuatrimestre_id=1')
            self.assertEqual({r['codigo']: r['docentes_sugeridos'] for r in criteria['abrir']}, {'c.1': 3})
            self.assertEqual([r['codigo'] for r in criteria['asincronica']], ['c.2'])
            self.assertEqual({r['codigo']: r['docentes_sugeridos'] for r in catalog}, {'c.1': 3, 'c.2': 0})
            checklist = self.get('/api/checklist-cierre?cuatrimestre_id=1')
            checks = checklist['controles']
            missing = next(r for r in checks if r['id'] == 'sin_docente')
            self.assertEqual(missing['cantidad'], 1)
            self.assertIn('15', missing['titulo'])
            response = self.client.get('/api/exportar/horarios?cuatrimestre_id=1')
            self.assertEqual(response.status_code, 200)
            workbook = load_workbook(io.BytesIO(response.content))
            rows = list(workbook['Criterio de Decisión'].values)[1:]
            self.assertEqual(next(r[5] for r in rows if r[1] == 'c.1'), 3)
            self.assertTrue(next(r[4] for r in rows if r[1] == 'c.2').startswith('ASINCRÓNICA'))

    def test_custom_duration_changes_fallback_but_keeps_explicit_end(self):
        sql("INSERT INTO asignaciones(catedra_id,cuatrimestre_id,docente_id,dia,hora_inicio,modalidad) VALUES (1,1,1,'Lunes','18:00','presencial')")
        with self.profile('institucion-ejemplo.json'):
            self.assertFalse(m.rangos_se_pisan('18:00', None, '19:00', '20:00'))
            self.assertTrue(m.rangos_se_pisan('18:00', '19:30', '19:00', '20:00'))
            report = self.get('/api/docentes/carga-horaria?cuatrimestre_id=1')
            self.assertEqual(report['docentes'][0]['horas'], 1.0)
            self.assertEqual(report['docentes'][0]['detalle'][0]['minutos'], 60)

    def control_file(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Titulo'] * 11)
        sheet.append(['N', 'Nombre', 'Apellido', 'DNI', 'Fecha', 'Inicio', '', '', '', 'Sede', 'Curso'])
        sheet.append([1, 'Alumno', 'Ficticio', 98000001, '01/03/2026', 'Marzo', '', '', '', 'Caballito', 'TECNICATURA DE PRUEBA (CABALLITO)'])
        data = io.BytesIO()
        workbook.save(data)
        return {'file': ('prueba.xlsx', data.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}

    def test_new_career_and_selected_year_in_control_and_export(self):
        sql("INSERT INTO plan_carrera(sede,carrera,anno,codigo_catedra,nombre_catedra) VALUES ('CABALLITO','TECNICATURA DE PRUEBA','2DO AÑO','c.1','Materia de prueba')")
        response = self.client.post('/api/control-inscripciones?cuatrimestre_id=2', files=self.control_file())
        self.assertEqual(response.status_code, 200)
        row = response.json()['results'][0]
        self.assertEqual(row['anno'], '2DO AÑO')
        self.assertNotEqual(row['estado'], 'SIN_PLAN')
        response = self.client.post('/api/control-inscripciones/exportar?cuatrimestre_id=2', files=self.control_file())
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.cell(2, 5).value, '2DO AÑO')

    def test_control_rejects_unknown_period(self):
        for route in ('/api/control-inscripciones', '/api/control-inscripciones/exportar'):
            response = self.client.post(route + '?cuatrimestre_id=999', files=self.control_file())
            self.assertEqual(response.status_code, 404)

    def test_custom_startup_does_not_load_iea_people_or_catalogs(self):
        sql('TRUNCATE docentes, catedras, cursos, sedes, cuatrimestres RESTART IDENTITY CASCADE')
        with self.profile('institucion-ejemplo.json'):
            asyncio.run(m.startup())
        for table in ('docentes', 'catedras', 'cursos', 'sedes'):
            self.assertEqual(sql('SELECT count(*) FROM ' + table)[0][0], 0)
        self.assertEqual(sql('SELECT count(*) FROM cuatrimestres')[0][0], 10)


if __name__ == '__main__':
    unittest.main(verbosity=2)
