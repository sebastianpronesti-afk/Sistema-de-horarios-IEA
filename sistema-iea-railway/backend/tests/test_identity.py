"""Identity migration and API regressions; exclusively disposable synthetic data."""
import io
import os
import unittest
from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import IntegrityError
import test_institution_api as base
import test_academic_workflow as workflow
from app.identity import document_key, code_key
from app.identity_migration import install_identity_guards


class IdentityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        workflow.AcademicWorkflowTests.setUpClass(); cls.client=workflow.AcademicWorkflowTests.client

    def setUp(self):
        workflow.AcademicWorkflowTests.setUp(self)
        base.sql("SELECT setval(pg_get_serial_sequence('docentes','id'),(SELECT max(id) FROM docentes))")
        base.sql("SELECT setval(pg_get_serial_sequence('catedras','id'),(SELECT max(id) FROM catedras))")

    edit=workflow.AcademicWorkflowTests.edit
    offer=workflow.AcademicWorkflowTests.offer

    def workbook(self, rows):
        wb=Workbook();sheet=wb.active
        for row in rows: sheet.append(row)
        data=io.BytesIO();wb.save(data)
        return {'file':('prueba.xlsx',data.getvalue())}

    def test_document_normalization_and_existing_id_survive_changes(self):
        for value in ('99.000.001','99 000 001','99000001.0',99000001):
            self.assertEqual(document_key(value),'99000001')
            response=self.client.post('/api/docentes',json={'nombre':'Otro','apellido':'Nombre','dni':value})
            self.assertEqual(response.status_code,409,response.text)
        response=self.client.put('/api/docentes/1/ficha',json={'nombre':'Nombre corregido','dni':'99.000.001'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['identificador'],'DOC-000001')
        self.assertEqual(base.sql('SELECT docente_id FROM docente_sede')[0][0],1)

    def test_old_update_route_cannot_bypass_document_or_id_checks(self):
        response=self.client.post('/api/docentes',json={'nombre':'Otra','apellido':'Persona','dni':'99700002'})
        self.assertEqual(response.status_code,200,response.text)
        for path in ('/api/docentes/2','/api/docentes/2/ficha'):
            self.assertEqual(self.client.put(path,json={'dni':'99.000.001'}).status_code,409)
            self.assertEqual(self.client.put(path,json={'id':999}).status_code,422)
        self.assertEqual(base.sql('SELECT dni FROM docentes WHERE id=2')[0][0],'99700002')

    def test_missing_document_names_require_review_and_homonyms_remain_distinct(self):
        first=self.client.post('/api/docentes',json={'nombre':'María','apellido':'Prueba'})
        self.assertEqual(first.status_code,200,first.text)
        self.assertEqual(self.client.post('/api/docentes',json={'nombre':' MARIA ','apellido':'PRUEBA'}).status_code,409)
        # Must complete the first fiche; a new document is not permission to duplicate it.
        self.assertEqual(self.client.post('/api/docentes',json={'nombre':'Maria','apellido':'Prueba','dni':'99700002'}).status_code,409)
        self.assertEqual(self.client.put('/api/docentes/2/ficha',json={'dni':'99700002'}).status_code,200)
        self.assertEqual(self.client.post('/api/docentes',json={'nombre':'Maria','apellido':'Prueba','dni':'99700003'}).status_code,200)
        report=self.client.get('/api/identidades/revision').json()
        self.assertIn([2,3],report['nombres_coincidentes'])
        self.assertEqual(report['documentos_repetidos'],[])

    def test_teacher_import_id_is_authoritative_and_repeated_import_is_idempotent(self):
        file=self.workbook([['docente_id','dni','nombre','apellido'],['DOC-000001','99.000.001','Nombre nuevo','Ficticio']])
        for _ in range(2):
            response=self.client.post('/api/importar/docentes',files=file)
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(response.json()['creados'],0)
        self.assertEqual(base.sql('SELECT id,nombre FROM docentes'),[(1,'Nombre nuevo')])
        conflict=self.workbook([['docente_id','dni','nombre'],['DOC-000001','99700002','Incorrecto']])
        self.assertEqual(self.client.post('/api/importar/docentes',files=conflict).status_code,422)
        self.assertEqual(base.sql('SELECT nombre FROM docentes')[0][0],'Nombre nuevo')

    def test_teacher_import_invalid_row_rolls_back_and_decimal_dni_is_correct(self):
        file=self.workbook([['dni','nombre','apellido'],[99700002.0,'Docente','Nuevo'],['NO-DNI','Dato','Inválido']])
        response=self.client.post('/api/importar/docentes',files=file)
        self.assertEqual(response.status_code,422,response.text)
        self.assertEqual(base.sql('SELECT count(*) FROM docentes')[0][0],1)
        file=self.workbook([['dni','nombre'],['99700002.0','Docente nuevo'],['99.700.002','Docente nuevo']])
        response=self.client.post('/api/importar/docentes',files=file)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(base.sql("SELECT dni FROM docentes WHERE id<>1"),[('99700002',)])

    def test_teacher_export_roundtrip_preserves_identifiers(self):
        response=self.client.get('/api/exportar/catalogo-docentes')
        self.assertEqual(response.status_code,200)
        sheet=load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.cell(2,1).value,'DOC-000001')
        result=self.client.post('/api/importar/docentes',files={'file':('docentes.xlsx',response.content)})
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()['creados'],0)

    def test_cuit_import_uses_same_document_identity(self):
        file=self.workbook([['CUIT','Apellido y nombre'],['20-99000001-0','Ficticio, Docente']])
        response=self.client.post('/api/importar/docentes-cuit',files=file)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['nuevos'],0)

    def test_database_guard_covers_raw_insert_and_code_punctuation_is_preserved(self):
        for query in ("INSERT INTO docentes(dni,nombre,apellido) VALUES ('99.000.001','X','Y')",
                      "INSERT INTO catedras(codigo,nombre) VALUES (' C.1 ','Otra')",
                      "UPDATE catedras SET codigo='c.999' WHERE id=1",
                      "UPDATE docentes SET id=999 WHERE id=1"):
            with self.assertRaises(IntegrityError): base.sql(query)
        base.sql("INSERT INTO catedras(codigo,nombre) VALUES ('MAT-1','Una'),('MAT1','Otra')")
        self.assertNotEqual(code_key('MAT-1'),code_key('MAT1'))
        self.assertEqual(base.sql('SELECT codigo FROM catedras WHERE id=1')[0][0],'c.1')

    def test_migration_keeps_legacy_duplicates_and_relations_then_blocks_growth(self):
        base.sql('ALTER TABLE docentes DISABLE TRIGGER identity_guard')
        try:
            base.sql("INSERT INTO docentes(id,dni,nombre,apellido) VALUES (5,'99.000.001','Registro','Anterior')")
        finally: base.sql('ALTER TABLE docentes ENABLE TRIGGER identity_guard')
        before=base.sql('SELECT * FROM docentes ORDER BY id')
        for _ in range(2):
            with base.SessionLocal() as db: install_identity_guards(db)
        self.assertEqual(before,base.sql('SELECT * FROM docentes ORDER BY id'))
        self.assertEqual(self.client.get('/api/identidades/revision').json()['documentos_repetidos'],[[1,5]])
        base.sql("UPDATE docentes SET nombre='Corregido' WHERE id=5")
        response=self.client.put('/api/docentes/5/ficha',json={'nombre':'Corrección por ficha','dni':'99.000.001'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(base.sql('SELECT dni FROM docentes WHERE id=5')[0][0],'99.000.001')
        with self.assertRaises(IntegrityError):
            base.sql("INSERT INTO docentes(dni,nombre,apellido) VALUES ('99 000 001','Otra','Ficha')")
        self.assertEqual(base.sql('SELECT docente_id FROM docente_sede')[0][0],1)

    def test_document_change_releases_only_previous_claim(self):
        base.sql("UPDATE docentes SET dni='99700002' WHERE id=1")
        base.sql("INSERT INTO docentes(dni,nombre,apellido) VALUES ('99000001','Otra','Persona')")
        with self.assertRaises(IntegrityError):
            base.sql("INSERT INTO docentes(dni,nombre,apellido) VALUES ('99.700.002','Otra','Persona')")

    @unittest.skipUnless(os.environ.get('IDENTITY_CONCURRENCY_TESTS')=='yes', 'Concurrent connections require native PostgreSQL')
    def test_concurrent_document_inserts_cannot_both_succeed(self):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        from sqlalchemy import text
        barrier=Barrier(2)
        def insert(document):
            with base.SessionLocal() as db:
                db.execute(text("SET LOCAL lock_timeout = '3s'"))
                db.execute(text("SET LOCAL statement_timeout = '5s'"))
                barrier.wait(timeout=5)
                try:
                    db.execute(text("INSERT INTO docentes(dni,nombre,apellido) VALUES (:dni,'Concurrente','Prueba')"),{'dni':document})
                    db.commit();return 'created'
                except IntegrityError:
                    db.rollback();return 'duplicate'
        with ThreadPoolExecutor(max_workers=2) as executor:
            a=executor.submit(insert,'99.700.088');b=executor.submit(insert,'99700088')
            self.assertCountEqual([a.result(timeout=10),b.result(timeout=10)],['created','duplicate'])
        self.assertEqual(base.sql("SELECT count(*) FROM docentes WHERE identity_document(dni)='99700088'")[0][0],1)

    def test_code_import_resolves_existing_code_without_changing_spelling_or_id(self):
        file=self.workbook([['Código','Materia'],[1,'C.1 Nombre corregido'],[1,'c.1 Nombre corregido']])
        response=self.client.post('/api/importar/catedras',files=file)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['creadas'],0)
        self.assertEqual(base.sql('SELECT codigo,nombre FROM catedras WHERE id=1')[0],('c.1','Nombre corregido'))
        self.assertEqual(self.client.put('/api/catedras/1',json={'codigo':'c.55'}).status_code,422)

    def test_generic_chair_codes_keep_punctuation_and_leading_zeros(self):
        file=self.workbook([['Código','Materia'],['MAT-01','Materia A'],['MAT01','Materia B'],['c.001','Materia C']])
        for _ in range(2):
            response=self.client.post('/api/importar/catedras',files=file)
            self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(base.sql("SELECT codigo FROM catedras WHERE codigo IN ('MAT-01','MAT01','c.001') ORDER BY codigo"),[('MAT-01',),('MAT01',),('c.001',)])

    def test_plan_identity_is_stable_and_complete_duplicate_is_rejected(self):
        fields={'resolucion':'RES 123/26','jurisdiccion':'CABA','modalidad':'presencial','version_plan':'1'}
        self.assertEqual(self.edit(operation='plan',plan_id='p1',revision=0,changes=fields).status_code,200)
        self.offer()
        conflict=self.edit(operation='plan',plan_id='p2',revision=1,changes={**fields,'etiqueta':'Otro nombre'})
        self.assertEqual(conflict.status_code,409,conflict.text)
        response=self.edit(operation='plan',plan_id='p1',revision=1,changes={'etiqueta':'Nombre corregido'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.client.get('/api/planes-estudio/p1').json()['identidad']['id'],'p1')
        view=self.client.get('/api/planificacion/1/planes/p1').json()
        self.assertTrue(view['materias'][0]['ofertada'])
        distinct=self.edit(operation='plan',plan_id='p2',revision=2,changes={**fields,'version_plan':'2'})
        self.assertEqual(distinct.status_code,200,distinct.text)

    def test_different_jurisdiction_is_a_different_plan_and_ids_cannot_be_edited(self):
        fields={'resolucion':'RES 123/26','jurisdiccion':'CABA','modalidad':'presencial'}
        self.assertEqual(self.edit(operation='plan',plan_id='p1',revision=0,changes=fields).status_code,200)
        self.assertEqual(self.edit(operation='plan',plan_id='p2',revision=1,changes={**fields,'jurisdiccion':'Provincia'}).status_code,200)
        self.assertEqual(self.edit(operation='plan',plan_id='p1',revision=2,changes={'id':'new-id'}).status_code,422)

    def test_repeated_new_plan_reuses_empty_draft_and_missing_metadata_is_visible(self):
        first=self.edit(operation='new_plan',career_id='career',revision=0).json()
        second=self.edit(operation='new_plan',career_id='career',revision=first['revision']).json()
        self.assertEqual(first['plan_id'],second['plan_id'])
        self.assertTrue(second['existente'])
        detail=self.client.get('/api/planes-estudio/'+first['plan_id']).json()
        self.assertEqual(detail['identidad']['estado'],'incompleta')
        self.assertEqual(detail['id'],first['plan_id'])


if __name__=='__main__': unittest.main()
