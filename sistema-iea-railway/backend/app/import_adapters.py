"""Pure spreadsheet adapters. They never receive a database session."""
import csv
from datetime import date, datetime, time
from decimal import Decimal
import io
import re
import unicodedata
from openpyxl import load_workbook


def norm(value):
    value = str(value or '').strip().lower()
    return re.sub(r'[^a-z0-9]', '', ''.join(c for c in unicodedata.normalize('NFD', value)
                                         if unicodedata.category(c) != 'Mn'))


FIELDS = {
    'asignacion_id': ['asignacion_id', 'id asignacion'],
    'codigo_materia': ['codigo_materia', 'codigo', 'codigo catedra', 'catedra'],
    'materia': ['materia', 'nombre materia', 'nombre catedra'],
    'carrera': ['carrera', 'curso', 'curso nombre'],
    'sede': ['sede', 'campus', 'sede referencia'],
    'periodo_id': ['periodo_id', 'cuatrimestre_id'],
    'comision': ['comision', 'grupo'], 'turno': ['turno'],
    'dia': ['dia', 'dia semana'], 'hora_inicio': ['hora inicio', 'hora', 'inicio'],
    'hora_fin': ['hora fin', 'fin'], 'docente': ['docente', 'profesor'],
    'docente_id': ['docente_id', 'id docente'], 'docente_dni': ['docente_dni', 'dni docente'],
    'modalidad': ['modalidad'], 'link_meet': ['link meet', 'link_meet', 'meet'],
    'recibe_alumnos_presenciales': ['recibe alumnos presenciales', 'recibe_alumnos_presenciales'],
    'dni': ['dni', 'documento', 'documento alumno'], 'alumno': ['alumno', 'apellido y nombre'],
    'nombre': ['nombre', 'nombres'], 'apellido': ['apellido', 'apellidos'],
    'email': ['email', 'correo'], 'es_edi': ['es_edi'], 'edi_materia': ['edi_materia'],
    'anno': ['anno', 'año', 'anio', 'año carrera'],
    'dia_tm': ['dia_tm', 'dia mañana'], 'hora_tm': ['hora_tm', 'hora mañana'],
    'dia_tn': ['dia_tn', 'dia noche'], 'hora_tn': ['hora_tn', 'hora noche'],
}
ALIASES = {norm(alias): field for field, aliases in FIELDS.items() for alias in aliases}


def cell_value(value):
    if isinstance(value, time): return value.strftime('%H:%M')
    if isinstance(value, (datetime, date)): return value.isoformat()
    return '' if value is None else str(value).strip()


def dni(value):
    """Preserve integer documents without appending the decimal zero from Excel."""
    value = cell_value(value).replace(' ', '').replace('-', '')
    if re.fullmatch(r'\d+\.0+', value): value = value.split('.')[0]
    else: value = value.replace('.', '')
    if not re.fullmatch(r'\d{6,12}', value):
        raise ValueError('Documento inválido; se requieren entre 6 y 12 dígitos')
    return value.lstrip('0') or '0'


def code(value):
    value = cell_value(value)
    if not value: raise ValueError('Falta el código de materia')
    if re.fullmatch(r'\d+(\.0+)?', value): return 'c.' + str(int(Decimal(value)))
    match = re.match(r'^(c\.\d+)(?:\s|$)', value, re.I)
    return match.group(1).lower() if match else value


def clock(value):
    value = cell_value(value)
    if not value: return None
    value = re.sub(r'\s*(hs|h)\s*$', '', value, flags=re.I).strip().replace('.', ':')
    match = re.fullmatch(r'(\d{1,2}):(\d{2})(?::00)?', value)
    if not match or int(match[1]) > 23 or int(match[2]) > 59:
        raise ValueError('Hora inválida; usá HH:MM entre 00:00 y 23:59')
    return f'{int(match[1]):02d}:{int(match[2]):02d}'


def read_sheets(content, filename):
    if len(content) > 15 * 1024 * 1024: raise ValueError('El archivo supera 15 MB')
    if filename.lower().endswith('.csv'):
        text = content.decode('utf-8-sig')
        try: dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
        except csv.Error: dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        if len(rows) > 15000: raise ValueError('El archivo supera 15.000 filas')
        return [('Datos', rows)]
    try: workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc: raise ValueError('No se pudo abrir el archivo XLSX o CSV') from exc
    try:
        sheets = []
        total = 0
        for sheet in workbook:
            if norm(sheet.title) in ('instructivo', 'instrucciones', 'configuracion'): continue
            rows = []
            for row in sheet.iter_rows(values_only=True):
                total += 1
                if total > 15000: raise ValueError('El archivo supera 15.000 filas')
                rows.append(list(row))
            sheets.append((sheet.title, rows))
        return sheets
    finally: workbook.close()


def parse_file(content, filename, kind, formato='auto', mapping=None):
    sheets = read_sheets(content, filename)
    records, errors, warnings, headers = [], [], [], []
    mapping = mapping or {}
    if not isinstance(mapping, dict) or set(mapping) - set(FIELDS):
        raise ValueError('El mapeo contiene campos desconocidos')
    if any(not isinstance(v, str) for v in mapping.values()): raise ValueError('Mapeo inválido')
    if len(set(norm(v) for v in mapping.values() if v)) != len([v for v in mapping.values() if v]):
        raise ValueError('Una columna no puede representar dos campos')
    for sheet, rows in sheets:
        if not any(any(cell_value(v) for v in row) for row in rows): continue
        headers.extend(cell_value(v) for v in rows[0] if cell_value(v))
        header_idx, indexes = None, {}
        for idx, row in enumerate(rows[:12]):
            candidate = {ALIASES[norm(value)]: pos for pos, value in enumerate(row) if norm(value) in ALIASES}
            if mapping:
                source = {norm(value): pos for pos, value in enumerate(row) if cell_value(value)}
                candidate.update({field: source[norm(value)] for field, value in mapping.items()
                                  if value and norm(value) in source})
            needed = {'codigo_materia', 'dni'} if kind == 'inscripciones' else {'codigo_materia', 'carrera', 'anno'} if kind == 'plan' else {'codigo_materia', 'dia', 'hora_inicio'}
            if needed.issubset(candidate):
                header_idx, indexes = idx, candidate
                headers.extend(cell_value(v) for v in row if cell_value(v))
                break
        if formato == 'iea' and kind == 'plan' and not mapping: header_idx = None
        if header_idx is not None:
            for index, row in enumerate(rows[header_idx + 1:], header_idx + 2):
                if not any(cell_value(v) for v in row): continue
                values = {field: cell_value(row[pos]) if pos < len(row) else '' for field, pos in indexes.items()}
                if any(v.startswith('=') for v in values.values()):
                    errors.append({'hoja': sheet, 'fila': index, 'mensaje': 'Hay fórmulas en campos de datos; pegá sus valores'})
                    continue
                values.update(_hoja=sheet, _fila=index, _formato='estandar')
                records.append(values)
        elif formato != 'estandar' and not mapping and kind == 'plan':
            parsed, issues = parse_iea_plan(sheet, rows)
            records.extend(parsed); errors.extend(issues)
        elif formato != 'estandar' and not mapping and kind == 'inscripciones':
            parsed, issues = parse_iea_enrollments(sheet, rows)
            records.extend(parsed); errors.extend(issues)
        elif formato != 'estandar' and not mapping and kind == 'horarios':
            # Classic IEA: code, matter, day, start, campus, teacher, link.
            count = 0
            for index, row in enumerate(rows, 1):
                if not row or not re.fullmatch(r'\d+(\.0+)?', cell_value(row[0])): continue
                vals = list(row) + [''] * 8
                records.append(dict(zip(['codigo_materia','materia','dia','hora_inicio','sede','docente','link_meet'], map(cell_value, vals[:7])),
                                    _hoja=sheet, _fila=index, _formato='iea'))
                count += 1
            if not count: errors.append({'hoja':sheet,'fila':1,'mensaje':'No se reconocieron las columnas del archivo'})
        else:
            headers.extend(cell_value(v) for v in rows[0] if cell_value(v))
            errors.append({'hoja':sheet,'fila':1,'mensaje':'Faltan columnas requeridas. Usá la plantilla o indicá el mapeo de columnas.'})
    if not records and not errors: errors.append({'hoja':'','fila':0,'mensaje':'El archivo no contiene registros. No se aplicará un reemplazo vacío.'})
    return records, errors, warnings, sorted(set(headers))


def parse_iea_enrollments(sheet, rows):
    records, errors = [], []
    codes = {m.group(1).lower() for row in rows for m in [re.match(r'^(c\.\d+)', cell_value(row[3]) if len(row)>3 else '', re.I)] if m}
    for index, row in enumerate(rows[1:], 2):
        if not any(cell_value(v) for v in row): continue
        vals = list(row) + [''] * 5
        student, document, matter, course = map(cell_value, vals[1:5])
        if not student and not document: continue
        match = re.match(r'^(c\.\d+)', matter, re.I)
        edi = 'EDI' in matter.upper() and not match
        if not match and not (edi and len(codes) == 1):
            errors.append({'hoja':sheet,'fila':index,'mensaje':'Materia sin código o EDI sin una única materia de referencia'})
            continue
        campus = re.search(r'\(([^)]+)\)', course)
        virtual = any(part in course.upper() for part in ('CIED', 'ONLINE', 'INTERIOR'))
        records.append({'codigo_materia':match.group(1).lower() if match else next(iter(codes)),
                        'dni':document,'alumno':re.sub(r'\s*\(\d+\).*$', '', student),
                        'carrera':course, 'sede': campus.group(1).strip() if campus else ('Online - Interior' if virtual else ''),
                        'modalidad':'virtual' if virtual else 'presencial',
                        'turno':'Mañana' if 'MAÑANA' in matter.upper() else 'Noche' if 'NOCHE' in matter.upper() else 'Virtual',
                        'es_edi':'true' if edi else 'false','edi_materia':matter if edi else '',
                        '_hoja':sheet,'_fila':index,'_formato':'iea'})
    return records, errors


def parse_iea_plan(sheet, rows):
    records, errors, pending = [], [], []
    career, year, reset = '', '', False
    edi_counts = {}
    def flush(selected_year):
        nonlocal pending
        for item in pending:
            item['anno'] = selected_year
            records.append(item)
        pending = []
    for index, row in enumerate(rows, 1):
        vals = list(row) + [''] * 11
        b, c, d, e = map(cell_value, vals[1:5])
        for value in (b, c, d):
            if ('TECNICO' in value.upper() or 'TECNICATURA' in value.upper()) and len(value) > 15:
                flush(year); career, year = value, ''
        if 'INSCRIPCION' in norm(c).upper(): year = ''; continue
        if 'PRACTICA FORMATIVA' in e.upper(): reset = True; continue
        if norm(d) == 'codigo' and norm(e) == 'materia': year = ''; continue
        if re.search(r'\d', c) and ('AÑO' in c.upper() or 'ANO' in c.upper()):
            flush(c); year = c
        try: subject = code(d) if re.fullmatch(r'\d+(\.0+)?', d) else None
        except ValueError: subject = None
        if subject and e:
            if not career:
                errors.append({'hoja':sheet,'fila':index,'mensaje':'Materia sin carrera identificable en la planilla IEA'}); continue
            professional = 'PROFESIONALIZANTE' in e.upper()
            if reset and not professional and year: year, reset = '', False
            item = {'sede':sheet,'carrera':career,'anno':year,'codigo_materia':subject,'materia':e,
                    'dia_tm':cell_value(vals[6]),'hora_tm':cell_value(vals[7]),
                    'dia_tn':cell_value(vals[9]),'hora_tn':cell_value(vals[10]),
                    '_hoja':sheet,'_fila':index,'_formato':'iea'}
            if year:
                records.append(item)
                if professional: year, reset = '', False
            else: pending.append(item)
        elif e.upper() == 'EDI' and career and year:
            edi_counts[career] = edi_counts.get(career, 0) + 1
            number = edi_counts[career]
            records.append({'sede':sheet,'carrera':career,'anno':year,'codigo_materia':f'EDI-{number}',
                            'materia':f'EDI {number} (Espacio de Definición Institucional)',
                            '_hoja':sheet,'_fila':index,'_formato':'iea'})
    flush(year)
    if not records and not errors: errors.append({'hoja':sheet,'fila':1,'mensaje':'No se reconoció un plan de carrera válido'})
    return records, errors
