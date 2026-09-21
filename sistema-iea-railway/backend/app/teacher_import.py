"""Teacher imports share the same ID/document rules as manual editing."""
import io
import re
from fastapi import HTTPException
from openpyxl import load_workbook
from app.identity import find_teacher, teacher_values, lock_teachers, name_key
from app.models.models import Docente


def import_teachers(db, content, cuit=False):
    if len(content) > 15 * 1024 * 1024:
        raise HTTPException(422, 'El archivo supera 15 MB')
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise HTTPException(422, 'No se pudo abrir el archivo XLSX de docentes') from exc
    created = updated = 0
    try:
        lock_teachers(db)
        for sheet in (workbook.worksheets if cuit else [workbook.worksheets[0]]):
            rows = sheet.iter_rows(values_only=True)
            if cuit:
                headers = None
            else:
                headers = [name_key(v) for v in next(rows, ())]
            for number, row in enumerate(rows, 1 if cuit else 2):
                if not any(v is not None and str(v).strip() for v in row): continue
                values = [str(v).strip() if v is not None else '' for v in row]
                if cuit:
                    if len(values) < 2 or ',' not in values[1]: continue
                    raw = re.sub(r'[-.\s]', '', values[0])
                    # Excel floats must not add a zero to the CUIT.
                    if re.fullmatch(r'[0-9]+\.0+', values[0]): raw = values[0].split('.')[0]
                    if not re.fullmatch(r'[0-9]{11}', raw):
                        raise ValueError(f'{sheet.title}, fila {number}: CUIT inválido')
                    surname, first = values[1].split(',', 1)
                    data = {'dni': raw[2:10], 'apellido': surname.strip(), 'nombre': first.strip()}
                else:
                    data = {}
                    for header, value in zip(headers, values):
                        if not value: continue
                        if header in ('docente id', 'id docente', 'id'): data['docente_id'] = value
                        elif header in ('dni', 'documento', 'dni docente'): data['dni'] = value
                        elif header in ('nombre', 'nombres'): data['nombre'] = value
                        elif header in ('apellido', 'apellidos'): data['apellido'] = value
                        elif header in ('email', 'mail', 'correo'): data['email'] = value
                        elif header in ('apellido y nombre', 'apellido nombre'):
                            if ',' not in value: raise ValueError(f'Fila {number}: separá apellido y nombre con una coma')
                            surname, first = value.split(',', 1)
                            data.update(apellido=surname.strip(), nombre=first.strip())
                    if not data: raise ValueError(f'Fila {number}: faltan columnas ID docente o documento y nombre/apellido')
                try:
                    existing = find_teacher(db, data)
                    clean = teacher_values(db, data, existing)
                except (ValueError, HTTPException) as exc:
                    message = exc.detail if isinstance(exc, HTTPException) else str(exc)
                    raise ValueError(f'{sheet.title}, fila {number}: {message}') from exc
                if existing:
                    for key, value in clean.items(): setattr(existing, key, value)
                    updated += 1
                else:
                    db.add(Docente(**clean)); created += 1
                db.flush()  # Repeated rows resolve to the same ID within this file.
        if not created and not updated:
            raise ValueError('El archivo no contiene docentes válidos')
        db.commit()
        return {'creados': created, 'nuevos': created, 'actualizados': updated, 'errores': []}
    except Exception as exc:
        db.rollback()
        if isinstance(exc, ValueError): raise HTTPException(422, str(exc)) from exc
        raise
    finally:
        workbook.close()
