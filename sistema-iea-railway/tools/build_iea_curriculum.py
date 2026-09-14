"""Convert the supplied IEA workbook into a reviewable, non-destructive catalog.

Reads XLSX only. Does not connect to a database or author/modify a workbook.
Usage: python tools/build_iea_curriculum.py source.xlsx output.json
"""
import argparse
from collections import Counter
from datetime import date, datetime
import hashlib
import json
import re
import unicodedata
import uuid
from pathlib import Path
from openpyxl import load_workbook


CONTROL = {'Instructivo', 'CONTROL EN SISTEMA', 'DOBLE TITULACION', 'Informacion a FEDE', 'control'}
NAMES = {'BANCARIA':'Administración Bancaria', 'SEGUROS':'Seguros',
         'NEGOCIOS DIGITALES':'Negocios Digitales', 'BIOSEGURIDAD':'Bioseguridad',
         'Acompañante terapeutico':'Acompañamiento Terapéutico', 'psicopedagogia':'Psicopedagogía',
         'Trabajo social':'Trabajo Social', 'Administracion':'Administración', 'Couseling':'Counseling',
         'Comercio':'Comercio Internacional', 'Despacho':'Despacho Aduanero y Régimen Aduanero',
         'Hoteleria':'Hotelería', 'Gestoria':'Gestoría', 'Guia de Turismo':'Guía de Turismo',
         'Relaciones Publicas':'Relaciones Públicas', 'Eventos':'Organización de Eventos',
         'ciencia de datos':'Ciencia de Datos e Inteligencia Artificial',
         'Administracion Agropecuaria':'Administración Agropecuaria', 'Gastronomia':'Gastronomía'}


def clean(value):
    if value is None: return ''
    if isinstance(value, (date, datetime)): return value.isoformat()
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return re.sub(r'\s+', ' ', str(value).replace('\\n', ' ')).strip()


def norm(value):
    return re.sub(r'[^a-z0-9]', '', ''.join(c for c in unicodedata.normalize('NFD', clean(value).lower())
                                         if unicodedata.category(c) != 'Mn'))


def uid(*parts):
    return str(uuid.uuid5(uuid.NAMESPACE_URL, 'iea:curriculum:' + ':'.join(map(str, parts))))


def number(value):
    match = re.fullmatch(r'([1-9])(?:\s*(?:°|º|er|do|ro|to|año|ano|cuatrimestre))*', clean(value), re.I)
    return int(match[1]) if match else None


def code(value):
    text = clean(value)
    if re.fullmatch(r'\d+(?:\.0+)?', text): return 'c.' + str(int(float(text)))
    match = re.fullmatch(r'c\.?\s*(\d+)', text, re.I)
    return 'c.' + str(int(match[1])) if match else None


def source(sheet, row, cell=None):
    result = {'hoja': sheet, 'fila': row}
    if cell: result['celda'] = cell
    return result


HEADERS = {'materias':'nombre', 'nombrecatedra':'nombre_sistema_archivo',
           'numerocatedra':'codigo_archivo', 'anoalcualcorresponde':'anio',
           'anoenelsistema':'anio_sistema_archivo', 'cuatrimestre':'cuatrimestre',
           'cuatrimestreenelsistema':'cuatrimestre_sistema_archivo',
           'cuatrimestreensistema':'cuatrimestre_sistema_archivo',
           'correlatividad':'correlatividades', 'correlatividadenelsistema':'correlatividades_sistema_archivo'}


def read_rows(sheet):
    # Propagate actual merged cells, never carry a plan/year through arbitrary blanks.
    merged = {}
    for area in sheet.merged_cells.ranges:
        if area.max_row > area.min_row:
            value = sheet.cell(area.min_row, area.min_col).value
            for row in range(area.min_row, area.max_row + 1): merged[row, area.min_col] = value
    max_row,max_col=sheet.max_row,sheet.max_column
    return [(r, {c: sheet.cell(r,c).value if sheet.cell(r,c).value is not None else merged.get((r,c))
                 for c in range(1, max_col + 1)}) for r in range(1, max_row + 1)]


def student_header(values):
    return {field: col for col, value in values.items() if (field := HEADERS.get(norm(value)))}


def make_subject(sheet, rownum, values, fields, plan_id, occurrences):
    raw = {key: clean(values.get(col)) for key,col in fields.items()}
    # Several sheets put the numeric code under the adjacent name header.
    if not raw.get('codigo_archivo') and code(raw.get('nombre_sistema_archivo')):
        raw['codigo_archivo'] = raw.pop('nombre_sistema_archivo')
    name = raw.get('nombre','')
    if not name or norm(name) == 'materias': return None
    if re.fullmatch(r'plan\s*\d+',name,re.I): return None
    if name.startswith('='): return None
    # Row headings (year/semester, resolutions) are not subjects.
    evidence = any(raw.get(k) for k in ('codigo_archivo','anio','cuatrimestre','anio_sistema_archivo'))
    if not evidence and not norm(name).startswith('edi'): return None
    issues = []
    original_name=name
    if re.fullmatch(r'\d+(?:\.0+)?',name):
        if not raw.get('codigo_archivo'): return None
        name='Nombre académico pendiente'
        issues.append('La denominación académica está guardada como número en el archivo')
    yr, sem = number(raw.get('anio')), number(raw.get('cuatrimestre'))
    if sem not in (1,2): sem = None
    if yr is None: issues.append('Año académico sin completar')
    if sem is None: issues.append('Cuatrimestre académico sin completar')
    for key in ('anio','cuatrimestre','correlatividades'):
        old = raw.get(key+'_sistema_archivo')
        if old and raw.get(key) and norm(old) != norm(raw[key]):
            if key == 'anio' and number(old) == yr: continue
            if key == 'cuatrimestre' and number(old) == sem: continue
            issues.append('Diferencia con '+key.replace('_',' ')+' del sistema anterior')
    source_code = code(raw.get('codigo_archivo'))
    if not source_code: issues.append('Código del archivo sin identificar')
    if any(v.startswith('=') for v in raw.values()): issues.append('Fórmula en un campo académico: revisar')
    key = (plan_id, norm(name), source_code, yr, sem)
    occurrences[key] += 1
    return {'id':uid(*key,occurrences[key]), 'nombre':name,'nombre_original':original_name,
            'codigo_archivo':source_code, 'codigo_archivo_original':raw.get('codigo_archivo',''),
            'nombre_sistema_archivo':raw.get('nombre_sistema_archivo',''),
            'anio':yr,'cuatrimestre':sem,'anio_original':raw.get('anio',''),
            'cuatrimestre_original':raw.get('cuatrimestre',''),
            'anio_sistema_archivo':raw.get('anio_sistema_archivo',''),
            'cuatrimestre_sistema_archivo':raw.get('cuatrimestre_sistema_archivo',''),
            'correlatividades':raw.get('correlatividades',''),
            'correlatividades_sistema_archivo':raw.get('correlatividades_sistema_archivo',''),
            'observaciones':issues, 'origen':source(sheet,rownum)}


def metadata(sheet, rows):
    header_row, columns = next(((r, {norm(v): c for c,v in vals.items() if clean(v)})
                               for r, vals in rows if any(norm(v)=='nombreoficial' for v in vals.values())), (0,{}))
    names = columns.get('nombreoficial'); resolution = columns.get('resolucion')
    title = columns.get('tituloqueseotorga')
    first_subject = next((r for r,v in rows if 'nombre' in student_header(v)), 5 if sheet=='BANCARIA' else 0)
    plans = []
    for r, vals in rows:
        if not names or r <= header_row or (first_subject and r >= first_subject): continue
        label = next((clean(v) for c,v in vals.items() if c < names and re.fullmatch(r'plan\s*\d+', clean(v), re.I)), '')
        if not label: continue
        official = clean(vals.get(names)); res_raw = vals.get(resolution); res = clean(res_raw)
        ambito = next((clean(v) for c,v in vals.items() if c < names and norm(v) in ('caba','capital','provincia','pcia','cied')), '')
        # Keep placeholders only if they have actual subjects later.
        observations = []
        if isinstance(res_raw, (date, datetime)):
            observations.append('Resolución guardada como fecha en Excel; verificar el número')
            res = ''
        if not norm(res): res=''
        if not official or norm(official) in ('nohay','nohayenprovincia') or set(official) <= {'-'}: official = ''
        if 'igualar' in norm(official) or 'tienequefigurar' in norm(official):
            observations.append(official); official = ''
        extra = ' · '.join(clean(v) for c,v in vals.items() if c > (title or names+2) and clean(v))
        status = 'por_confirmar'
        if re.search(r'\b(viejo|anterior)\b', extra, re.I): status = 'anterior'
        if re.search(r'\b(no se da|no va)\b', extra, re.I): status = 'no_ofertado'
        if re.search(r'\bactivo\b', extra, re.I): status = 'vigente_informado'
        starts = re.search(r'(?:desde(?: el)?|1\s*año)\s*(20\d\d)', extra, re.I)
        start_year = int(starts[1]) if starts else None
        if start_year and start_year > 2026: status = 'futuro'
        if not official: observations.append('Nombre oficial pendiente')
        if not res: observations.append('Resolución pendiente')
        if norm(ambito)=='cied': modality, jurisdiction = 'distancia', None
        elif norm(ambito) in ('caba','capital','provincia','pcia'):
            modality, jurisdiction = 'presencial', 'CABA' if norm(ambito) in ('caba','capital') else 'Provincia de Buenos Aires'
        else: modality, jurisdiction = None, None
        if jurisdiction is None: observations.append('Jurisdicción aprobante pendiente')
        career_id = uid('carrera',sheet)
        academic_title=clean(vals.get(title))
        if not academic_title:
            academic_title=next((clean(v) for c,v in vals.items() if c>(resolution or names) and 'tecnic' in norm(v)), '')
        plans.append({'id':uid('plan',career_id,norm(label),ambito,res or clean(res_raw)),
                      'carrera_id':career_id,'etiqueta':label.title(),'nombre_oficial':official,
                      'resolucion':res,'resolucion_original':clean(res_raw),
                      'titulo':academic_title, 'ambito_origen':ambito,
                      'modalidad':modality,'jurisdiccion':jurisdiction,'situacion':status,
                      'inicio_informado':start_year,'nota_vigencia':extra,
                      'observaciones':observations,'origen':source(sheet,r),'materias':[]})
    return plans


def tertiary(sheet, rows):
    plans = metadata(sheet,rows); by_label = {norm(p['etiqueta']):p for p in plans}
    fields = {}; occurrences = Counter(); pending = []; block_plan=None; pending_header=False
    meaningful = [p for p in plans if p['nombre_oficial'] or p['resolucion']]
    for r, vals in rows:
        # A proposal embedded after the approved plan must not be appended to it.
        if any('presentaciondedistancia' in norm(v) for v in vals.values()):
            career_id=uid('carrera',sheet);block_plan='propuestadistancia'
            proposal={'id':uid('plan',career_id,block_plan),'carrera_id':career_id,'etiqueta':'Propuesta a distancia',
                      'nombre_oficial':'','resolucion':'','resolucion_original':'','titulo':'','ambito_origen':'Presentación de distancia',
                      'modalidad':'distancia','jurisdiccion':None,'situacion':'por_confirmar','inicio_informado':None,
                      'nota_vigencia':'Presentación de distancia: aprobación no informada',
                      'observaciones':['Propuesta sin resolución aprobatoria informada'],'origen':source(sheet,r),'materias':[]}
            plans.append(proposal);by_label[block_plan]=proposal;fields={};pending_header=True;continue
        labels=[norm(v) for v in vals.values() if re.fullmatch(r'plan\s*\d+',clean(v),re.I)]
        known_resolution=bool(labels and labels[0] in by_label and norm(by_label[labels[0]]['resolucion']) and
                              any(norm(v)==norm(by_label[labels[0]]['resolucion']) for v in vals.values()))
        if labels and (known_resolution or any('tecnic' in norm(v) for v in vals.values())) and any(norm(v) in ('caba','capital','provincia','pcia','cied') for v in vals.values()):
            block_plan=labels[0] if fields else None
            pending_header=True
            continue
        header = student_header(vals)
        if 'nombre' in header:
            if not pending_header: block_plan=None
            pending_header=False
            fields = header
            continue
        if sheet == 'BANCARIA' and r >= 5:
            fields = {'nombre':3,'codigo_archivo':4,'anio':6,'cuatrimestre':7}
        if not fields: continue
        # Desarrollo Humano has subject names in D and sequential numbers in C.
        if sheet == 'Desarrollo Humano':
            fields = {**fields, 'nombre':4}; fields.pop('nombre_sistema_archivo',None)
        label = next((norm(v) for c,v in vals.items() if c < fields['nombre'] and re.fullmatch(r'plan\s*\d+',clean(v),re.I)), None)
        conflict=bool(label and block_plan and label!=block_plan)
        plan = None if conflict else by_label.get(label or block_plan) if (label or block_plan) else meaningful[0] if len(meaningful)==1 else None
        if label and not plan and not conflict:
            career_id=uid('carrera',sheet)
            plan={'id':uid('plan',career_id,label),'carrera_id':career_id,'etiqueta':label.replace('plan','Plan '),
                  'nombre_oficial':'','resolucion':'','resolucion_original':'','titulo':'','ambito_origen':'',
                  'modalidad':None,'jurisdiccion':None,'situacion':'por_confirmar','inicio_informado':None,
                  'nota_vigencia':'','observaciones':['Cabecera del plan sin identificar'],
                  'origen':source(sheet,r),'materias':[]}
            plans.append(plan);by_label[label]=plan
        subject = make_subject(sheet,r,vals,fields,plan['id'] if plan else uid('pendientes',sheet),occurrences)
        if subject:
            if conflict:
                subject['observaciones'].append(f'La fila dice {label}, pero la cabecera del bloque indica {block_plan}')
            if plan: plan['materias'].append(subject)
            else:
                subject['observaciones'].append('No se pudo determinar el plan de esta materia')
                pending.append(subject)
    # Empty form labels such as Plan 7 are not actual plans.
    plans = [p for p in plans if p['materias'] or p['nombre_oficial'] or p['resolucion']]
    for plan in plans:
        if not plan['materias']: plan['observaciones'].append('Plan sin detalle de materias')
    return {'id':uid('carrera',sheet),'nombre':NAMES.get(sheet,sheet),'nivel':'terciario',
            'origen':sheet,'planes':plans,'materias_sin_plan':pending}


def secondary(sheet, rows):
    ident=uid('carrera',sheet); pid=uid('plan',ident)
    plan={'id':pid,'carrera_id':ident,'etiqueta':sheet,'nombre_oficial':'Bachiller en Economía y Administración' if sheet=='BCE 2026' else 'BEA',
          'resolucion':'Dictamen N°3160/2025' if sheet=='BCE 2026' else '',
          'resolucion_original':'Dictamen N°3160/2025' if sheet=='BCE 2026' else '',
          'titulo':'','ambito_origen':'Provincia' if sheet=='BCE 2026' else '',
          'jurisdiccion':'Provincia de Buenos Aires' if sheet=='BCE 2026' else None,
          'modalidad':None,'situacion':'por_confirmar','inicio_informado':None,'nota_vigencia':'',
          'observaciones':['Modalidad pendiente'],'origen':source(sheet,2 if sheet=='BCE 2026' else 4),'materias':[]}
    occurrences=Counter(); modules=[]
    for r,vals in rows:
        name=clean(vals.get(3)); sem=number(vals.get(4))
        if 'módulo de estudio' in name.lower():
            modules.append({'nombre':name,'codigo_archivo':code(vals.get(2)),'cuatrimestre':sem,'origen':source(sheet,r)})
            continue
        match=re.match(r'^(c\.\d+)\s+(.+)$',name,re.I)
        source_code=match[1] if match else code(vals.get(2))
        if not source_code or not name: continue
        values={1:match[2] if match else name,2:source_code,3:sem}
        item=make_subject(sheet,r,values,{'nombre':1,'codigo_archivo':2,'cuatrimestre':3},pid,occurrences)
        # Modules are preserved separately; no academic year is inferred from a code.
        plan['materias'].append(item)
    plan['modulos']=modules
    return {'id':ident,'nombre':sheet,'nivel':'secundario','origen':sheet,'planes':[plan],'materias_sin_plan':[]}


def combination(sheet, rows):
    # Comparison worksheets are evidence for articulations, not new standalone degrees.
    matrices=[]; current=[]; occurrences=Counter(); ident=uid('articulacion',sheet)
    for r, vals in rows:
        if any(re.fullmatch(r'plan\s*\d+',clean(v),re.I) for v in vals.values()) and any(norm(v) in ('caba','capital','provincia','pcia','cied') for v in vals.values()):
            continue
        starts=[c for c,v in vals.items() if norm(v)=='materias']
        if starts:
            current=[]
            for idx,start in enumerate(starts):
                stop=starts[idx+1] if idx+1<len(starts) else max(vals)+1
                fields=student_header({c:v for c,v in vals.items() if start<=c<stop})
                group={'nombre':f'Bloque {len(matrices)+1}','origen':source(sheet,r),'materias':[]}
                matrices.append(group);current.append((fields,group))
            continue
        if not current and sheet=='Doble titulacion comercio  mas ':
            group={'nombre':'Comparación de materias','origen':source(sheet,1),'materias':[]};matrices.append(group)
            current=[({'nombre':3,'nombre_sistema_archivo':4,'codigo_archivo':5,'anio':6,'anio_sistema_archivo':7,
                      'cuatrimestre':8,'cuatrimestre_sistema_archivo':9,'correlatividades':10},group)]
        for fields,group in current:
            subject=make_subject(sheet,r,vals,fields,ident,occurrences)
            if subject:
                subject['plan_origen']=next((clean(vals.get(c)) for c in range(max(1,fields['nombre']-2),fields['nombre'])
                                            if re.fullmatch(r'plan\s*\d+',clean(vals.get(c)),re.I)),'')
                group['materias'].append(subject)
    return {'id':ident,'nombre':sheet.strip(),'situacion':'por_confirmar',
            'observaciones':['Confirmar los planes participantes y las equivalencias antes de aplicar esta articulación'],
            'origen':source(sheet,1),'bloques':matrices}


def convert(path):
    book=load_workbook(path,data_only=False)
    careers=[];combinations=[]; ignored=[]
    for sheet in book:
        if sheet.title in CONTROL: ignored.append(sheet.title);continue
        rows=read_rows(sheet)
        if sheet.title=='Doble Titulacion Faltantes':
            for r,vals in rows:
                if clean(vals.get(2)):
                    combinations.append({'id':uid('articulacion',sheet.title,clean(vals[2])),
                        'nombre':clean(vals[2]),'situacion':'por_confirmar','origen':source(sheet.title,r),
                        'observaciones':['Sin detalle de materias', clean(vals.get(3)),clean(vals.get(4))],'bloques':[]})
        elif 'doble' in norm(sheet.title) or sheet.title=='Publicidad y Marketing':
            combinations.append(combination(sheet.title,rows))
        elif sheet.title in ('BEA','BCE 2026'): careers.append(secondary(sheet.title,rows))
        else: careers.append(tertiary(sheet.title,rows))
    book.close()
    return {'schema_version':1,'institution_id':'iea','source':{'filename':Path(path).name,
            'sha256':hashlib.sha256(Path(path).read_bytes()).hexdigest(),'revision':'2026-09-14','excluded_control_sheets':ignored},
            'careers':careers,'articulations':combinations}


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('source');parser.add_argument('output');args=parser.parse_args()
    result=convert(args.source)
    Path(args.output).parent.mkdir(parents=True,exist_ok=True)
    Path(args.output).write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'careers':len(result['careers']),'plans':sum(len(c['planes']) for c in result['careers']),
          'subjects':sum(len(p['materias']) for c in result['careers'] for p in c['planes']),
          'unassigned':sum(len(c['materias_sin_plan']) for c in result['careers']),
          'articulations':len(result['articulations'])}))
